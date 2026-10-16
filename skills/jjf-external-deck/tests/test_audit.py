import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT=Path(__file__).resolve().parents[1]
SCRIPT=ROOT/'scripts/audit_pptx.py'
class AuditTests(unittest.TestCase):
    def setUp(self):
        self.assertTrue(SCRIPT.is_file(),'RED: editable-structure audit is not implemented')
        spec=importlib.util.spec_from_file_location('audit_pptx',SCRIPT)
        self.m=importlib.util.module_from_spec(spec);spec.loader.exec_module(self.m)
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.d=Path(self.tmp.name)
    def deck(self,text='社区增长引擎',background=False):
        p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
        s=p.slides.add_slide(p.slide_layouts[6]);h=None
        if background:
            path=self.d/'source.png';Image.new('RGB',(1920,1080),(2,4,3)).save(path)
            h=hashlib.sha256(path.read_bytes()).hexdigest()
            s.shapes.add_picture(str(path),0,0,p.slide_width,p.slide_height)
        if text:
            s.shapes.add_textbox(Inches(1),Inches(1),Inches(10),Inches(1)).text=text
        f=self.d/'deck.pptx';p.save(f);return f,h
    def test_native_text_passes_structure_only(self):
        f,_=self.deck();r=self.m.audit(f)
        self.assertTrue(r['structure_passed']);self.assertFalse(r['release_ready'])
    def test_flat_raster_fails(self):
        f,_=self.deck(text=None,background=True);r=self.m.audit(f)
        self.assertFalse(r['structure_passed'])
    def test_full_source_plus_overlay_fails(self):
        f,h=self.deck(background=True)
        r=self.m.audit(f,{'pages':[{'source_page_id':'p006','source_sha256':h,'visible_text':['社区增长引擎']}]})
        self.assertTrue(any('SOURCE_RASTER_OVERLAY' in e for e in r['errors']))
    def test_missing_text_fails_parity(self):
        f,_=self.deck();r=self.m.audit(f,{'pages':[{'source_page_id':'p006','source_sha256':'no','visible_text':['联合交付引擎']}]})
        self.assertFalse(r['content_parity_passed'])
    def test_mapping_and_parity(self):
        f,_=self.deck();r=self.m.audit(f,{'pages':[{'source_page_id':'p006','source_sha256':'no','visible_text':['社区增长引擎']}]})
        self.assertTrue(r['content_parity_passed']);self.assertEqual(r['pages'][0]['source_page_id'],'p006')
        self.assertFalse(r['release_ready'])
    def test_invalid_zip(self):
        f=self.d/'bad.pptx';f.write_text('not zip');r=self.m.audit(f)
        self.assertFalse(r['structure_passed'])
if __name__=='__main__': unittest.main()
